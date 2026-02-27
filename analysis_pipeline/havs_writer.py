"""
HAVS Template Writer
=====================
Writes detected events to XLSX in the customer-specified HAVS proforma format.

Layout matches the Altair HAVS proforma:
  A: Date         — dd/mm/yyyy, shown once per date change
  B: Session      — session number(s) for this event
  C: Time         — HH:MM-HH:MM event time range
  D: Spectrogram Primary Frequencies (Hz) — one row per frequency (up to 5)
  E: Spectrogram Harmonics (Hz)           — detected harmonic multiples, or 'Not Present'
  F: LOFAR Primary Frequencies (Hz)       — mirrors column D
  G: LOFAR Harmonics (Hz)                 — mirrors column E
  H: DEMON Blade Rate (Hz)
  I: DEMON Shaft Rate (Hz)
  J: Event classification
  K: Remarks

Each event uses N_FREQS_PER_EVENT rows (default 5), one per detected tonal frequency.
Harmonic detection checks whether 2×, 3×, or higher multiples of each frequency
appear among the other detected tonals for that event.

© Oravont Systems LLP — Confidential
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from datetime import datetime

# =====================================================================
# CONFIGURATION
# =====================================================================
N_FREQS_PER_EVENT = 5       # number of tonal frequency rows per event
HARMONIC_TOL_HZ = 2.0       # tolerance for matching harmonics (±Hz)
MAX_HARMONIC_MULT = 4       # check up to Nx harmonics

# =====================================================================
# STYLES (matching customer proforma exactly)
# =====================================================================
_HEADER_FONT = Font(bold=True, size=10, name='Arial', color='FFFFFF')
_HEADER_FILL = PatternFill('solid', fgColor='4472C4')
_SUB_HEADER_FILL = PatternFill('solid', fgColor='D6E4F0')
_SUB_HEADER_FONT = Font(bold=True, size=9, name='Arial')
_DATA_FONT = Font(size=9, name='Arial')
_DATE_FONT = Font(size=9, name='Arial', bold=True)
_THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'),
)
_CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
_LEFT = Alignment(horizontal='left', vertical='center', wrap_text=True)


# =====================================================================
# HARMONIC DETECTION
# =====================================================================

def _find_harmonics(primary_hz, all_freqs_hz, tol=HARMONIC_TOL_HZ):
    """Check if multiples of primary_hz appear in all_freqs_hz.

    Checks 2×, 3×, ... up to MAX_HARMONIC_MULT× of primary_hz.
    Returns comma-separated string of found harmonics, or 'Not Present'.
    """
    found = []
    for mult in range(2, MAX_HARMONIC_MULT + 1):
        target = primary_hz * mult
        for f in all_freqs_hz:
            if abs(f - target) <= tol and f != primary_hz:
                found.append(int(round(f)))
                break
    if found:
        return ', '.join(str(h) for h in found)
    return 'Not Present'


# =====================================================================
# MAIN WRITER
# =====================================================================

def write_havs_xlsx(events, output_path, notes=None):
    """Write events to HAVS proforma XLSX matching customer template.

    Parameters:
        events: list of event dicts from event_detector.
                Expected keys: start, end, assessment, remarks,
                blade_rate, shaft_rate, top_tonal_freqs,
                session (optional), all_tonal_freqs (optional).
        output_path: path to write .xlsx
        notes: optional list of note strings for footer
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'HAVS Results'

    # ─── Row 1: Main headers ───
    headers_r1 = {
        'A1': 'Date',
        'B1': 'Session',
        'C1': 'Time',
        'D1': 'Spectrogram Features',
        'F1': 'LOFAR Features',
        'H1': 'DEMON Features',
        'J1': 'Event',
        'K1': 'Remarks',
    }
    for cell_ref, text in headers_r1.items():
        cell = ws[cell_ref]
        cell.value = text
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _THIN_BORDER

    # Merge header cells for multi-column groups
    ws.merge_cells('D1:E1')   # Spectrogram Features
    ws.merge_cells('F1:G1')   # LOFAR Features
    ws.merge_cells('H1:I1')   # DEMON Features

    # Fill background for all cells in row 1
    for col in range(1, 12):
        cell = ws.cell(row=1, column=col)
        cell.border = _THIN_BORDER
        if not cell.value:
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT

    # ─── Row 2: Sub-headers ───
    sub_headers = [
        None, None, None,                          # A, B, C
        'Primary Frequencies (Hz)',                 # D
        'Harmonics (Hz)',                           # E
        'Primary Frequencies (Hz)',                 # F
        'Harmonics (Hz)',                           # G
        'Blade Rate (Hz)',                          # H
        'Shaft Rate (Hz)',                          # I
        None, None,                                 # J, K
    ]
    for col, text in enumerate(sub_headers, 1):
        cell = ws.cell(row=2, column=col)
        if text:
            cell.value = text
        cell.font = _SUB_HEADER_FONT
        cell.fill = _SUB_HEADER_FILL
        cell.alignment = _CENTER
        cell.border = _THIN_BORDER

    # ─── Column widths ───
    widths = {
        'A': 12,   # Date
        'B': 10,   # Session
        'C': 14,   # Time
        'D': 22,   # Spectrogram Primary
        'E': 22,   # Spectrogram Harmonics
        'F': 22,   # LOFAR Primary
        'G': 22,   # LOFAR Harmonics
        'H': 14,   # Blade Rate
        'I': 14,   # Shaft Rate
        'J': 20,   # Event
        'K': 55,   # Remarks
    }
    for letter, w in widths.items():
        ws.column_dimensions[letter].width = w

    # ─── Write events ───
    row = 3
    current_date = None

    for ev in sorted(events, key=lambda e: e['start']):
        date_str = ev['start'].strftime('%d/%m/%Y')
        time_str = (f"{ev['start'].strftime('%H:%M')}-"
                    f"{ev['end'].strftime('%H:%M')}")

        # Collect tonal frequencies for this event
        # Prefer all_tonal_freqs (full list), fall back to top_tonal_freqs
        all_freqs = ev.get('all_tonal_freqs', ev.get('top_tonal_freqs', []))
        top_freqs = list(all_freqs[:N_FREQS_PER_EVENT])

        # Pad to N_FREQS_PER_EVENT
        while len(top_freqs) < N_FREQS_PER_EVENT:
            top_freqs.append(None)

        # Session label
        session_str = ''
        if ev.get('sessions'):
            sessions = sorted(ev['sessions'])
            if len(sessions) == 1:
                session_str = str(sessions[0])
            else:
                session_str = f'{sessions[0]}-{sessions[-1]}'
        elif ev.get('session'):
            session_str = str(ev['session'])

        # All detected frequencies as floats for harmonic checking
        all_freq_values = [f for f in all_freqs if f is not None and f > 0]

        # Write N_FREQS_PER_EVENT rows
        for fi in range(N_FREQS_PER_EVENT):
            freq = top_freqs[fi]
            is_first_row = (fi == 0)

            # Column A: Date (first row of event, only on date change)
            if is_first_row and date_str != current_date:
                cell = ws.cell(row=row, column=1, value=date_str)
                cell.font = _DATE_FONT
                cell.alignment = _LEFT
                current_date = date_str

            # Column B: Session (first row only)
            if is_first_row and session_str:
                cell = ws.cell(row=row, column=2, value=session_str)
                cell.font = _DATA_FONT
                cell.alignment = _CENTER

            # Column C: Time (first row only)
            if is_first_row:
                cell = ws.cell(row=row, column=3, value=time_str)
                cell.font = _DATA_FONT
                cell.alignment = _LEFT

            # Columns D, E: Spectrogram features
            if freq is not None and freq > 0:
                freq_int = int(round(freq))
                ws.cell(row=row, column=4, value=freq_int).font = _DATA_FONT
                ws.cell(row=row, column=4).alignment = _CENTER

                harmonics = _find_harmonics(freq, all_freq_values)
                ws.cell(row=row, column=5, value=harmonics).font = _DATA_FONT
                ws.cell(row=row, column=5).alignment = _CENTER

                # Columns F, G: LOFAR features (mirror spectrogram)
                ws.cell(row=row, column=6, value=freq_int).font = _DATA_FONT
                ws.cell(row=row, column=6).alignment = _CENTER
                ws.cell(row=row, column=7, value=harmonics).font = _DATA_FONT
                ws.cell(row=row, column=7).alignment = _CENTER

            # Columns H, I: DEMON features (first row only)
            if is_first_row:
                br = ev.get('blade_rate', 0)
                sr = ev.get('shaft_rate', 0)
                if br and br > 0:
                    ws.cell(row=row, column=8, value=round(br, 1)).font = _DATA_FONT
                    ws.cell(row=row, column=8).alignment = _CENTER
                if sr and sr > 0:
                    ws.cell(row=row, column=9, value=round(sr, 2)).font = _DATA_FONT
                    ws.cell(row=row, column=9).alignment = _CENTER

            # Column J: Event (first row only)
            if is_first_row:
                ws.cell(row=row, column=10, value=ev.get('assessment', '')).font = _DATA_FONT
                ws.cell(row=row, column=10).alignment = _LEFT

            # Column K: Remarks (first row only)
            if is_first_row:
                ws.cell(row=row, column=11, value=ev.get('remarks', '')).font = _DATA_FONT
                ws.cell(row=row, column=11).alignment = _LEFT

            # Borders on all cells in this row
            for col in range(1, 12):
                ws.cell(row=row, column=col).border = _THIN_BORDER

            row += 1

    # ─── Footer notes ───
    if notes:
        row += 2
        ws.cell(row=row, column=1, value='Notes:').font = Font(
            size=8, bold=True, name='Arial')
        for note in notes:
            row += 1
            ws.cell(row=row, column=1, value=note).font = Font(
                size=8, italic=True, name='Arial')

    # Freeze panes below headers
    ws.freeze_panes = 'A3'

    wb.save(output_path)
    return output_path
